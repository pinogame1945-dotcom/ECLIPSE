"""Editable KEV *observed-diff* briefing (not dateAdded-based reporting)."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from eclipse_reports.daily import _base, _filter_row, _rect, _strip, _txt, C, HREF


def render_events(events_path: Path, profile_path: Path, output_path: Path,
                  *, max_details: int = 5) -> dict:
    payload = json.loads(events_path.read_text(encoding='utf-8'))
    config = json.loads(profile_path.read_text(encoding='utf-8'))
    if not isinstance(config.get('customer_name'), str) or not config['customer_name'].strip():
        raise ValueError('customer_name required')
    if not config.get('include_all') and not config.get('vendors') and not config.get('products'):
        raise ValueError('An explicit vendor/product policy or include_all is required')
    if payload.get('baseline') or not payload.get('events'):
        raise ValueError('Refusing to describe baseline/no-change as new alerts')
    events = [x for x in payload['events'] if x.get('material_change') and _filter_row(x['record'], config)]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = 'PROJECT ECLIPSE / Official KEV change detection'
    prs.core_properties.subject = 'First observed source changes; NOT exploit detection time'

    title = prs.slides.add_slide(prs.slide_layouts[6])
    title.background.fill.solid()
    title.background.fill.fore_color.rgb = C['navy']
    _rect(title, .68, .82, .7, .1, 'cyan')
    _txt(title, 'PROJECT ECLIPSE', .7, 1.2, 11.8, .52, size=23, color='cyan', bold=True)
    _txt(title, 'KEV CHANGE BRIEFING', .7, 2.0, 11.9, .82, size=38, color='white', bold=True)
    _txt(title, config['customer_name'], .7, 3.22, 11.7, .58, size=23, color='white')
    _txt(title, f"Observed by ECLIPSE  {payload['observed_utc']}", .7, 4.14, 11.8, .38,
         size=15, color='white')
    _txt(title, f'{len(events)} material changes matching this public sample profile', .7, 4.78, 11.8, .44,
         size=18, color='cyan')
    _txt(title, 'CISA KEV  |  source changes ≠ exploitation start', .7, 6.67, 11.8, .32,
         size=11, color='white')

    overview = _base(prs, '変更サマリー', f"ECLIPSE observed at {payload['observed_utc']}")
    counts = {kind: sum(e['kind'] == kind for e in events) for kind in ('added', 'updated', 'removed')}
    _txt(overview, f"新規観測: {counts['added']}件", .8, 2.17, 11.5, .5, size=24, color='red', bold=True)
    _txt(overview, f"既存項目の更新: {counts['updated']}件", .8, 2.85, 11.5, .48, size=22, color='blue', bold=True)
    _txt(overview, f"カタログからの削除: {counts['removed']}件", .8, 3.47, 11.5, .48, size=22, color='amber', bold=True)
    _rect(overview, .75, 4.28, 11.75, 1.35, 'white', 'line', radius=True)
    _txt(overview, 'KEV登録日（dateAdded）とECLIPSEの検知日時は異なる。\n削除は「悪用がない」ことを意味しない。対象資産との照合は未実施。',
         .96, 4.55, 11.25, .96, size=14)

    for page in range(max(1, (len(events)+6)//7)):
        subset = events[page*7:(page+1)*7]
        slide = _base(prs, '情報源の変更一覧', f"{config['customer_name']} / page {page+1}")
        _rect(slide, .72, 2.03, 11.88, .48, 'navy2')
        _txt(slide, '変化 / CVE', .89, 2.14, 2.68, .27, size=12, color='white', bold=True)
        _txt(slide, 'メーカー・製品', 3.75, 2.14, 3.23, .27, size=12, color='white', bold=True)
        _txt(slide, '変更点', 7.12, 2.14, 5.07, .27, size=12, color='white', bold=True)
        if not subset:
            _txt(slide, 'この監視条件に一致する重要な変更なし', .88, 3.0, 10.8, .6, size=20, color='muted')
        for i, event in enumerate(subset):
            row = event['record']
            y = 2.7 + .56*i
            _txt(slide, _strip(f"{event['kind'].upper()} / {event['cve_id']}", 40), .84, y, 2.83, .46, size=11, bold=True)
            _txt(slide, _strip(f"{row.get('vendorProject','')} / {row.get('product','')}", 50), 3.73, y, 3.18, .46, size=10)
            changes = ', '.join(event.get('changed_fields', [])) if event['kind'] == 'updated' else event['kind']
            _txt(slide, _strip(changes, 90), 7.09, y, 5.15, .45, size=10)

    for i, event in enumerate(events[:max_details], 1):
        row = event['record']
        slide = _base(prs, f"差分解説 {i:02d} / {event['cve_id']}",
                      f"{event['kind']}  /  {row.get('vendorProject', '-')} — {row.get('product', '-')}")
        _rect(slide, .7, 2.04, 11.86, 4.58, 'white', 'line', radius=True)
        _txt(slide, 'CISA説明（原文）', 1.02, 2.29, 10.8, .35, size=16, bold=True)
        _txt(slide, _strip(row.get('shortDescription', '説明なし'), 370), 1.02, 2.82, 10.5, 1.13, size=13)
        _txt(slide, 'CISA掲載の対応文', 1.02, 4.1, 10.8, .36, size=15, bold=True)
        _txt(slide, _strip(row.get('requiredAction', '未記載'), 300), 1.02, 4.57, 10.8, .8, size=12)
        update = ', '.join(event['changed_fields']) if event['kind']=='updated' else 'CISA KEV catalog membership'
        _txt(slide, f"変更項目: {_strip(update, 130)}\nKEV dateAdded: {row.get('dateAdded','-')}  |  ECLIPSE observed: {event['observed_utc']}",
             1.02, 5.72, 11, .66, size=10, color='muted')

    sources = _base(prs, '出典と解釈', 'SOURCE / REPRODUCIBILITY')
    _txt(sources, HREF + '\n' + payload['source'] + '\n' + 'Feed SHA-256: '+payload['source_sha256'],
         .85, 2.15, 11.6, 1.15, size=13)
    _txt(sources, 'このレポートはCISA KEVの項目追加・内容更新・削除を検知した結果。\n実際の攻撃発生日、顧客影響、緊急度を独自に認定するものではない。\nメーカー公式の影響条件と修正版は別途確認が必要。',
         .85, 4.2, 11.6, 1.7, size=15)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    manifest = {
        'observed_utc': payload['observed_utc'], 'kev_feed_sha256': payload['source_sha256'],
        'customer_name': config['customer_name'], 'selected_count': len(events),
        'selected_events': [e['event_id'] for e in events], 'counts': counts,
        'interpretation': 'Changes in CISA KEV catalog since last persisted observed state (not exploitation start).',
    }
    output_path.with_suffix('.manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument('--events', required=True, type=Path)
    parser.add_argument('--customer', required=True, type=Path)
    parser.add_argument('--out', required=True, type=Path)
    args = parser.parse_args()
    print(json.dumps(render_events(args.events, args.customer, args.out), ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
