"""Build a PPTX from an official CISA KEV JSON feed, with no speculative AI claims.

The KEV feed's dateAdded is a calendar date, not a disclosure or detection timestamp.
This tool does not claim zero vulnerabilities outside KEV, or assess customer assets.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime, date, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

C = {
    'navy': RGBColor(10, 23, 43), 'navy2': RGBColor(16, 33, 57),
    'ink': RGBColor(33, 48, 67), 'muted': RGBColor(104, 119, 139),
    'white': RGBColor(255, 255, 255), 'bg': RGBColor(247, 249, 252),
    'blue': RGBColor(43, 119, 239), 'cyan': RGBColor(44, 202, 209),
    'amber': RGBColor(247, 164, 52), 'red': RGBColor(218, 74, 83),
    'line': RGBColor(225, 232, 239), 'green': RGBColor(28, 151, 117),
}
FONTS = 'Noto Sans CJK JP'
HREF = 'https://www.cisa.gov/known-exploited-vulnerabilities-catalog'


def _txt(slide, text, x, y, w, h, *, size=14, color='ink', bold=False, align=None, valign=None, max_lines=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(.02); tf.margin_right = Inches(.02)
    tf.margin_top = Inches(.02); tf.margin_bottom = Inches(.02)
    if valign:
        tf.vertical_anchor = valign
    lines = str(text).splitlines() or ['']
    if max_lines is not None:
        lines = lines[:max_lines]
    for j, line in enumerate(lines):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONTS
        p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = C[color]
        p.space_before = Pt(0); p.space_after = Pt(5)
        if align is not None:
            p.alignment = align
    return shape


def _rect(slide, x, y, w, h, color='white', stroke=None, radius=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = C[color]
    sh.line.fill.background() if not stroke else None
    if stroke:
        sh.line.color.rgb = C[stroke]
    if radius:
        try: sh.adjustments[0] = 0.06
        except Exception: pass
    return sh


def _base(prs, title, tag, *, demo=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = C['bg']
    _rect(slide, 0, 0, 13.33, .14, 'cyan')
    _txt(slide, 'ECLIPSE  /  THREAT INTELLIGENCE', .55, .28, 8.5, .29, size=10, color='muted', bold=True)
    _txt(slide, title, .55, .84, 12.1, .65, size=27, bold=True)
    _txt(slide, tag, .58, 1.51, 12.1, .33, size=11, color='muted')
    _rect(slide, .55, 7.1, 12.2, .015, 'line')
    _txt(slide, 'SOURCE: CISA KEV  |  dateAdded = catalog registration date', .58, 7.16, 11.5, .22, size=8, color='muted')
    _txt(slide, f'{len(prs.slides):02d}', 12.21, 7.15, .5, .25, size=9, color='muted', align=PP_ALIGN.RIGHT)
    if demo:
        _txt(slide, 'DEMO • SYNTHETIC DATA', 9.65, .28, 3.1, .35, size=10, color='red', bold=True, align=PP_ALIGN.RIGHT)
    return slide


def _strip(raw: str, maxchars=340) -> str:
    clean = ' '.join((raw or '').split())
    if len(clean) > maxchars:
        return clean[:maxchars-1] + '…'
    return clean


def _filter_row(row, config):
    # Vendor and product filters are OR within each group, AND across groups.
    # If both are empty, an explicit include_all=true is required to prevent leakage.
    vendors = [str(s).casefold() for s in config.get('vendors', [])]
    products = [str(s).casefold() for s in config.get('products', [])]
    if not vendors and not products and not config.get('include_all', False):
        raise ValueError('No matching policy: specify vendors/products or explicit include_all=true')
    vendor = str(row.get('vendorProject') or '').casefold()
    product = str(row.get('product') or '').casefold()
    if vendors and not any(v == vendor or v in vendor for v in vendors):
        return False
    if products and not any(p == product or p in product for p in products):
        return False
    return True


def _select(feed, config, start, end):
    if not isinstance(feed, dict) or not isinstance(feed.get('vulnerabilities'), list):
        raise ValueError('Invalid KEV JSON: vulnerabilities[] not found')
    # Validate even when no KEV entries exist in the chosen date window.
    if not config.get('vendors') and not config.get('products') and not config.get('include_all', False):
        raise ValueError('No matching policy: specify vendors/products or explicit include_all=true')
    selected = []
    errors = []
    for r in feed['vulnerabilities']:
        if not isinstance(r, dict):
            errors.append('non-object row'); continue
        try: day = date.fromisoformat(r['dateAdded'])
        except (ValueError, KeyError, TypeError):
            errors.append(str(r.get('cveID','unknown'))); continue
        if start <= day <= end and _filter_row(r, config):
            selected.append(r)
    selected.sort(key=lambda x: (x['dateAdded'], x.get('vendorProject',''), x.get('cveID','')), reverse=True)
    # Malformed inputs must not be silently reported as complete.
    if errors:
        raise ValueError(f'Invalid KEV dateAdded records: {len(errors)} (e.g. {errors[0]})')
    return selected


def _add_title(prs, config, start, end, rows, generated, feed, demo):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = C['navy']
    _rect(slide, .55, .72, .65, .1, 'cyan')
    _txt(slide, 'PROJECT ECLIPSE', .65, 1.15, 11.6, .45, size=20, color='cyan', bold=True)
    _txt(slide, 'DAILY / PERIODIC\nTHREAT BRIEFING', .63, 1.92, 11.5, 1.65, size=38, color='white', bold=True)
    _rect(slide, .65, 4.25, 11.9, .025, 'cyan')
    _txt(slide, config['customer_name'], .66, 4.56, 11.5, .49, size=25, color='white', bold=True)
    _txt(slide, f'KEV registration window  {start.isoformat()} — {end.isoformat()}', .66, 5.25, 11.5, .33, size=15, color='white')
    _txt(slide, f'Published from official KEV JSON  •  {len(rows)} matching new entries', .66, 5.75, 11.2, .35, size=12, color='cyan')
    _txt(slide, f'Built at {generated.strftime("%Y-%m-%d %H:%M UTC")}', .66, 6.77, 8.2, .31, size=9, color='white')
    if demo:
        _txt(slide, 'SAMPLE ONLY · ALL CVE DATA IS FICTIONAL', 7.25, 6.68, 5.4, .36, size=10, color='amber', bold=True, align=PP_ALIGN.RIGHT)


def _card(slide, x, y, number, label, sub, *, accent='blue'):
    _rect(slide, x, y, 3.82, 1.58, 'white', 'line', radius=True)
    _rect(slide, x, y, .07, 1.58, accent)
    _txt(slide, str(number), x+.28, y+.16, 3.35, .64, size=32, color=accent, bold=True)
    _txt(slide, label, x+.3, y+.89, 3.34, .34, size=13, bold=True)
    _txt(slide, sub, x+.3, y+1.26, 3.35, .25, size=9, color='muted')


def _overview(prs, rows, start, end, config, demo):
    s = _base(prs, 'エグゼクティブサマリー', f'{start} – {end}  /  {config["customer_name"]}', demo=demo)
    vendors = sorted(set((r.get('vendorProject') or 'unknown') for r in rows))
    ransomware = sum(1 for r in rows if str(r.get('knownRansomwareCampaignUse', '')).lower() == 'known')
    _card(s, .62, 2.12, len(rows), '新規 KEV 該当', 'KEV dateAdded が対象期間内', accent='red')
    _card(s, 4.75, 2.12, len(vendors), '対象メーカー数', '顧客の監視設定に一致', accent='blue')
    _card(s, 8.88, 2.12, ransomware, 'ランサムウェア関連', 'CISAの項目が Known のみ', accent='amber')
    _rect(s, .62, 4.03, 12.02, 1.88, 'white', 'line', radius=True)
    _txt(s, '本日の判断ポイント', .9, 4.27, 11.4, .34, size=18, bold=True)
    if rows:
        msg = 'KEV掲載は実際の悪用が確認されたことを示す。対象製品の利用有無を確認し、メーカー公告の影響条件・修正版を別途照合する。'
    else:
        msg = '今回指定した監視条件と期間では、KEVに新規追加された項目は見つからなかった。CVE全体・メーカー公告などで新たな問題がないことは意味しない。'
    _txt(s, msg, .9, 4.9, 11.22, .9, size=15)
    _txt(s, f'対象メーカー: {", ".join(config.get("vendors", [])) or "条件なし"}  |  製品: {", ".join(config.get("products", [])) or "全製品"}', .68, 6.38, 11.9, .37, size=10, color='muted')


def _list_page(prs, rows, config, start, end, demo, page=0, page_size=8):
    subset = rows[page*page_size:(page+1)*page_size]
    s = _base(prs, '新規 KEV 一覧', f'{config["customer_name"]}  /  {start} – {end}  /  page {page+1}', demo=demo)
    columns = [(.7, 2.18, 'CVE'), (2.93, 1.52, '登録日'), (4.47, 2.65, 'メーカー / 製品'), (7.22, 5.3, '脆弱性の概要（CISA記載の要約）')]
    _rect(s, .68, 2.06, 12.0, .53, 'navy2', radius=True)
    for x, w, label in columns:
        _txt(s, label, x+.06, 2.16, w-.08, .28, size=11, color='white', bold=True)
    for i, r in enumerate(subset):
        y = 2.75 + i * .49
        _rect(s, .68, y-.035, 12, .47, 'white' if i % 2 == 0 else 'bg')
        _txt(s, r.get('cveID', '—'), .76, y, 2.1, .32, size=10, bold=True)
        _txt(s, r.get('dateAdded', '—'), 3.0, y, 1.38, .32, size=10)
        _txt(s, _strip(f'{r.get("vendorProject","")} / {r.get("product","")}', 33), 4.53, y, 2.58, .32, size=10)
        _txt(s, _strip(r.get('vulnerabilityName') or r.get('shortDescription', ''), 75), 7.24, y, 5.24, .41, size=9)
    if not subset:
        _txt(s, '該当する新規 KEV 登録はありません。', .81, 3.18, 11.4, .55, size=20, color='muted')


def _detail(prs, r, index, demo):
    s = _base(prs, f'重点解説  {index:02d} / {r.get("cveID", "-")}',
              f'{r.get("vendorProject", "-")}  /  {r.get("product", "-")}  /  KEV {r.get("dateAdded", "-")}', demo=demo)
    _rect(s, .68, 2.0, 7.5, 4.52, 'white', 'line', radius=True)
    _rect(s, 8.45, 2.0, 4.18, 4.52, 'navy2', radius=True)
    _txt(s, '何が起きるか — CISAの説明', .96, 2.28, 6.94, .33, size=17, bold=True)
    _txt(s, _strip(r.get('shortDescription', 'CISA説明未記載'), 510), .96, 2.83, 6.94, 1.66, size=14)
    _txt(s, 'なぜ重要か', .96, 4.55, 6.93, .34, size=17, bold=True)
    _txt(s, 'CISA KEVへの掲載は実際の悪用が確認されたことを示す。ただし自社環境の該当有無は、製品・バージョン・構成情報による別途確認が必要。', .96, 5.05, 6.91, 1.11, size=12)
    _txt(s, 'CISA 推奨対応（原文要約せず記載）', 8.75, 2.32, 3.49, .64, size=14, color='cyan', bold=True)
    _txt(s, _strip(r.get('requiredAction', 'Published guidance を参照'), 430), 8.75, 3.12, 3.5, 1.81, size=12, color='white')
    _txt(s, '別途確認', 8.75, 5.18, 3.38, .32, size=13, color='cyan', bold=True)
    _txt(s, '・メーカーの最新公告\n・影響バージョン\n・修正済みバージョン', 8.75, 5.52, 3.3, .86, size=10, color='white')
    _txt(s, f'CVE: {r.get("cveID", "-")}  |  Catalog: {HREF}', .74, 6.63, 11.7, .29, size=9, color='muted')


def _sources(prs, config, start, end, feed, digest, errors, demo):
    s = _base(prs, '出典・判定上の留意事項', f'{config["customer_name"]}  /  provenance & limitations', demo=demo)
    _rect(s, .7, 2.04, 11.92, 4.67, 'white', 'line', radius=True)
    notes = [
        '① 出典：米国CISA KEV公式JSON。本文はKEV登録情報に基づく。',
        '② dateAdded は KEV 登録日。脆弱性の公開日・攻撃開始日・自動検知時刻ではない。',
        '③ KEVに未掲載であっても、悪用なし・影響なしを意味しない。',
        '④ 対象製品・修正バージョンは、メーカー公式公告の追加確認が必要。',
        '⑤ メーカー/製品名のキーワード一致で分類。表記揺れによる漏れがあり得る。',
        '⑥ 顧客の資産台帳との照合、CVSS/EPSS、CVE以外の警告はこの試作の範囲外。',
    ]
    for i, n in enumerate(notes):
        _txt(s, n, .97, 2.33 + i*.6, 11.32, .46, size=13)
    _txt(s, f'KEV feed dateReleased: {feed.get("dateReleased", "unknown")}  /  catalogVersion: {feed.get("catalogVersion", "unknown")}', .97, 6.08, 11.0, .25, size=10, color='muted')
    _txt(s, f'SHA256: {digest}', .97, 6.36, 11.06, .25, size=9, color='muted')


def render(kev_path: Path, config_path: Path, out_path: Path, start: date, end: date,
           *, demo=False, max_details=3):
    if end < start:
        raise ValueError('End date precedes start date')
    raw = kev_path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    feed = json.loads(raw)
    if not demo:
        # A truncated or test feed must never silently produce a clean production report.
        if not isinstance(feed, dict) or not isinstance(feed.get('vulnerabilities'), list) or len(feed['vulnerabilities']) < 1000:
            raise ValueError('Production KEV feed too small or invalid (<1000 entries)')
        reported = feed.get('count')
        if reported is not None and reported != len(feed['vulnerabilities']):
            raise ValueError('KEV count does not match the supplied catalog')
    config = json.loads(config_path.read_text(encoding='utf-8'))
    if not isinstance(config.get('customer_name'), str) or not config['customer_name'].strip():
        raise ValueError('customer_name is required')
    rows = _select(feed, config, start, end)
    generated = datetime.now(timezone.utc)
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    prs.core_properties.title = f'ECLIPSE: {config["customer_name"]} KEV {start} to {end}'
    prs.core_properties.subject = 'CISA KEV dateAdded-based monitoring report'
    prs.core_properties.author = 'PROJECT ECLIPSE'
    _add_title(prs, config, start, end, rows, generated, feed, demo)
    _overview(prs, rows, start, end, config, demo)
    for i in range(max(1, (len(rows) + 7)//8)):
        _list_page(prs, rows, config, start, end, demo, i)
    for i, r in enumerate(rows[:max_details], 1):
        _detail(prs, r, i, demo)
    _sources(prs, config, start, end, feed, digest, [], demo)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    manifest = {
        'generated_utc': generated.isoformat(timespec='seconds'),
        'customer_name': config['customer_name'], 'period': {'start': start.isoformat(), 'end': end.isoformat()},
        'kev_feed_sha256': digest, 'kev_feed_date_released': feed.get('dateReleased'),
        'kev_catalog_version': feed.get('catalogVersion'),
        'selected': [r.get('cveID') for r in rows], 'selected_count': len(rows),
        'classification': 'KEV dateAdded registration; NOT observed-in-real-time detection',
        'synthetic_demo': demo,
    }
    manifest_path = out_path.with_suffix('.manifest.json')
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    return manifest


def main():
    p = argparse.ArgumentParser(description='Generate customer-filtered CISA KEV presentation')
    p.add_argument('--kev-json', required=True, type=Path)
    p.add_argument('--customer', required=True, type=Path)
    p.add_argument('--output', required=True, type=Path)
    p.add_argument('--date', type=date.fromisoformat, default=None,
                   help='dateAdded (US catalog date), YYYY-MM-DD; default yesterday JST')
    p.add_argument('--days', type=int, default=1, help='range length inclusive, e.g. 7')
    p.add_argument('--demo', action='store_true', help='watermark all slides SAMPLE')
    a = p.parse_args()
    if not 1 <= a.days <= 31:
        p.error('--days must be within 1..31')
    end = a.date or datetime.now(ZoneInfo('Asia/Tokyo')).date() - timedelta(days=1)
    start = end - timedelta(days=a.days-1)
    print(json.dumps(render(a.kev_json, a.customer, a.output, start, end, demo=a.demo), ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
