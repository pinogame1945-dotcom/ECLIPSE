"""Editable private PowerPoint for a validated customer-product CVE match report.

Never claims affected installed versions, confirmed exposure, patch readiness,
or real customer compromise. Caller must keep real inventory/report PRIVATE.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from eclipse_reports.daily import C,_rect,_strip,_txt


def _page(prs,heading,sub):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb=C["bg"]
    _rect(s,0,0,13.333,.14,"cyan")
    _txt(s,"ECLIPSE   /   CUSTOMER PRODUCT VULNERABILITY WATCH",.65,.34,11.6,.34,size=11,bold=True,color="muted")
    _txt(s,_strip(heading,100),.65,.88,11.95,.75,size=26,bold=True)
    _txt(s,_strip(sub,155),.69,1.69,11.8,.42,size=11,color="muted")
    _rect(s,.70,7.09,11.9,.025,"line")
    _txt(s,"CVE v5 • CISA KEV • FIRST EPSS  |  Installed-version impact: NOT ASSESSED",.72,7.14,11.6,.25,size=9,color="muted")
    _txt(s,f"{len(prs.slides):02d}",12.35,7.12,.28,.29,size=9,color="muted")
    return s


def _card(s,x,y,title,value,color="blue"):
    _rect(s,x,y,3.77,1.35,"white","line",radius=True)
    _rect(s,x,y,.075,1.35,color)
    _txt(s,_strip(str(value),23),x+.25,y+.2,3.20,.55,size=28,color=color,bold=True)
    _txt(s,_strip(title,30),x+.25,y+.84,3.15,.32,size=12)


def render_report(report_path:Path,output_path:Path,*,max_details:int=8,
                  max_list:int=40)->dict:
    if not 0<=max_details<=25 or not 0<=max_list<=100:
        raise ValueError("Invalid slide cap")
    raw=report_path.read_bytes()
    d=json.loads(raw)
    if (d.get("schema_version")!=1 or
            d.get("classification")!="CUSTOMER_PRODUCT_REFERENCE_REPORT_NOT_VERIFIED_ASSET_EXPOSURE"
            or not isinstance(d.get("products"),list) or not isinstance(d.get("preview"),list)):
        raise ValueError("Unsupported product-CVE report")
    prs=Presentation()
    prs.slide_width=Inches(13.333)
    prs.slide_height=Inches(7.5)
    prs.core_properties.title="ECLIPSE / Customer product vulnerability review"
    prs.core_properties.subject="Publisher product mentions; installed-version impact not assessed"
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=C["navy"]
    _rect(s,.76,1.14,.8,.12,"cyan")
    _txt(s,"PROJECT ECLIPSE",.78,1.47,11.8,.53,size=23,color="cyan",bold=True)
    _txt(s,"製品別 脆弱性情報レポート",.78,2.24,11.85,1.04,size=35,color="white",bold=True)
    _txt(s,_strip(d["customer_name"],85),.8,3.88,11.7,.64,size=24,color="white",bold=True)
    _txt(s,f"参照データ時点: {d['as_of']}  /  対象: {'過去全件' if d['mode']=='all' else d['lookback_start_inclusive']+'以降の新規公開・更新および全KEV'}",
         .80,4.75,11.55,.63,size=14,color="white")
    _txt(s,"顧客資産の影響・導入バージョンへの該当性は未判定",.81,5.91,11.65,.45,size=16,color="cyan",bold=True)
    if d["synthetic_demo"]:
        _txt(s,"SYNTHETIC CUSTOMER PROFILE • DEMONSTRATION ONLY",.81,6.67,11.65,.34,size=11,color="red",bold=True)

    n=d["summary"]
    s=_page(prs,"エグゼクティブサマリー","監視対象の製品名とCVE原本の発行元表記を厳密照合した結果")
    _card(s,.70,2.35,"登録された製品数",n["monitored_inventory_entries"])
    _card(s,4.80,2.35,"報告対象の製品-CVE件数",n["product_cve_rows_in_report"],"cyan")
    _card(s,8.91,2.35,"KEV掲載（製品-CVE件数）",n["kev_listed_product_cve_rows"],"red")
    _rect(s,.72,4.23,11.86,1.73,"white","line",radius=True)
    _txt(s,"重要：一致とは「脆弱性が存在する」と断定することではない",.97,4.50,11.14,.39,size=17,bold=True,color="red")
    _txt(s,"製品名がCVEの affected[] に掲載されている候補。対象バージョンが対象外である場合もある。"
         "導入バージョン・構成・適用済み修正・資産存在を確認するまで、顧客への影響は未判定。",
         .97,5.04,11.14,.81,size=13)
    _txt(s,f"全履歴の製品-CVE照合件数: {n['published_product_cve_pairs_all_history']:,}  /  "
           f"報告条件: {d['mode']}  /  一覧省略: {d['preview_omitted_rows']:,} 件（CSVに全件保持）",
         .75,6.3,11.7,.4,size=11,color="muted")

    products=d["products"]
    for offset in range(0,len(products),8):
        subset=products[offset:offset+8]
        s=_page(prs,"監視対象の製品一覧",f"登録件数: {len(products)}  /  {offset+1}〜{offset+len(subset)}件を表示")
        _rect(s,.77,2.18,11.75,.51,"navy2")
        _txt(s,"製品名 / 登録バージョン",.96,2.29,6.15,.32,size=12,color="white",bold=True)
        _txt(s,"報告件数",7.52,2.29,1.71,.32,size=12,color="white",bold=True)
        _txt(s,"KEV掲載",9.47,2.29,2.17,.32,size=12,color="white",bold=True)
        for i,p in enumerate(subset):
            y=2.81+i*.47
            _txt(s,_strip(f"{p['vendor']} / {p['product']}  [{p['installed_version'] or 'バージョン未登録'}]",91),
                 1.0,y,6.12,.37,size=11,bold=i==0)
            _txt(s,str(p["included_product_cve_rows"]),7.74,y,1.55,.37,size=11)
            _txt(s,str(p["kev_listed_included"]),9.65,y,1.43,.37,size=11,
                 color="red" if p["kev_listed_included"] else "muted")
        if len(subset)<8:
            _txt(s,"一致0件でも安全性を証明するものではありません。製品名の別表記や未収集の脆弱性があり得ます。",
                 1.0,6.74,11.5,.28,size=10,color="red")

    topics=d["preview"][:max_list]
    for offset in range(0,max(1,len(topics)),8):
        subset=topics[offset:offset+8]
        s=_page(prs,"注目するCVE候補",
                f"対象 {n['product_cve_rows_in_report']} 製品-CVE件中の主要候補を表示 / 全件はCSV参照")
        _rect(s,.74,2.15,11.85,.53,"navy2")
        _txt(s,"CVE番号",.90,2.26,2.25,.32,size=11,color="white",bold=True)
        _txt(s,"該当製品",3.22,2.26,3.88,.32,size=11,color="white",bold=True)
        _txt(s,"KEV / CVSS",7.20,2.26,1.64,.32,size=11,color="white",bold=True)
        _txt(s,"調査理由",9.01,2.26,3.10,.32,size=11,color="white",bold=True)
        if not subset:
            _txt(s,"抽出された候補はありません。脆弱性が存在しないという意味ではありません。",
                 1.04,3.02,11.08,.62,size=16,color="muted")
        for i,item in enumerate(subset):
            y=2.85+i*.46
            _txt(s,item["cve_id"],.9,y,2.27,.31,size=10,bold=True)
            _txt(s,_strip(f"{item['vendor']} / {item['product']}",50),3.23,y,3.8,.31,size=10)
            _txt(s,_strip(("KEV " if item["kev_in_snapshot"] else "")+
                 (str(item["cvss_score"]) if item["cvss_score"] is not None else "CVSS不明"),26),
                 7.20,y,1.74,.31,size=10,color="red" if item["kev_in_snapshot"] else "ink")
            _txt(s,_strip(item["review_queue"],40),9.02,y,3.3,.33,size=9)

    for i,item in enumerate(d["preview"][:max_details],1):
        s=_page(prs,"個別確認： "+item["cve_id"],
                f"{item['vendor']} / {item['product']}  |  導入版: {item['installed_version'] or '未登録'}")
        _rect(s,.76,2.17,7.63,4.59,"white","line",radius=True)
        _rect(s,8.56,2.17,3.99,4.59,"navy2",radius=True)
        _txt(s,"CVE原本の説明",.97,2.40,7.0,.38,size=16,bold=True)
        _txt(s,_strip(item["publisher_description"],390),.97,2.87,6.99,1.40,size=12)
        _txt(s,"出版社による影響バージョン記述（未判定）",.97,4.37,7.0,.38,size=14,bold=True)
        sources=item.get("publisher_version_statements") or []
        raw_versions=json.dumps([{"container":v["container"],"defaultStatus":v["default_status"],
             "versions":v["versions_raw"]} for v in sources[:3]],ensure_ascii=False)
        _txt(s,_strip(raw_versions,390),.97,4.87,6.99,1.37,size=11)
        _txt(s,"根拠と優先確認",8.88,2.45,3.33,.40,size=17,color="cyan",bold=True)
        _txt(s,f"KEV: {'掲載' if item['kev_in_snapshot'] else 'スナップショット未掲載'}",
             8.88,3.07,3.32,.42,size=14,color="white")
        _txt(s,f"CVSS: {item['cvss_score'] if item['cvss_score'] is not None else '不明'}",8.88,3.65,3.29,.40,size=13,color="white")
        val=item["epss_probability"]
        _txt(s,f"EPSS: {val:.4f}" if val is not None else "EPSS: 未取得",8.88,4.20,3.31,.41,size=13,color="white")
        _txt(s,f"EPSS鮮度: {item['epss_state']}",8.88,4.74,3.33,.35,size=11,color="white")
        _txt(s,"顧客への影響: 未判定",8.88,5.48,3.3,.45,size=13,color="cyan",bold=True)
        link=_txt(s,"一次情報: "+_strip(item["cve_reference"],112),.82,6.84,11.69,.29,size=10,color="blue")
        for p in link.text_frame.paragraphs:
            for run in p.runs:run.hyperlink.address=item["cve_reference"]

    s=_page(prs,"根拠・制限事項","アナリストによる確認が必要な事項")
    notes=[
        "照合範囲：CVE原本 affected[] のメーカー名と製品名の完全一致（大文字小文字は無視）。",
        "製品の導入・稼働実態、インストール版の影響、回避策・パッチ適用状況は確認していない。",
        "CVE原本: 2026-10-10固定。KEV/EPSS: 倉庫構築時点の別取得データ。",
        "KEV掲載=野外での悪用をCISAが確認。顧客での攻撃発生を示さない。",
        "EPSS=世界的な悪用の予測であり、顧客固有の被害確率ではない。",
        "詳細・原文の証拠と全件リストは、同時生成したJSON・CSVを参照。",
    ]
    for i,t in enumerate(notes):
        _rect(s,.81,2.17+i*.67,11.61,.55,"white","line",radius=True)
        _txt(s,t,1.03,2.29+i*.67,11.18,.38,size=12)
    if d["synthetic_demo"]:
        for slide in prs.slides:
            _txt(slide,"SYNTHETIC CUSTOMER ONLY",10.25,.24,2.9,.35,size=9,color="red",bold=True)

    output_path.parent.mkdir(parents=True,exist_ok=True)
    prs.save(str(output_path))
    manifest={"pptx_file":output_path.name,"slides":len(prs.slides),
              "input_report_sha256":hashlib.sha256(raw).hexdigest(),
              "reported_product_cve_rows":n["product_cve_rows_in_report"],
              "displayed_list_rows":len(topics),
              "displayed_detail_rows":min(max_details,len(d["preview"])),
              "synthetic_demo":bool(d["synthetic_demo"]),
              "customer_exposure":"NOT_ASSESSED"}
    output_path.with_suffix(".manifest.json").write_text(
        json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    return manifest


def main():
    p=argparse.ArgumentParser(description="Editable PPTX from offline private product vulnerability report")
    p.add_argument("--report-json",required=True,type=Path)
    p.add_argument("--out",required=True,type=Path)
    p.add_argument("--max-details",type=int,default=8)
    p.add_argument("--max-list",type=int,default=40)
    args=p.parse_args()
    os.umask(0o077)
    if args.report_json.resolve()==args.out.resolve():
        p.error("Output cannot overwrite input JSON")
    manifest=render_report(args.report_json,args.out,max_details=args.max_details,max_list=args.max_list)
    print(json.dumps(manifest,ensure_ascii=False))


if __name__=="__main__":main()
