"""Editable customer-profiled PPTX from evidence-locked multifeed briefing.

Does not require an LLM or disclose private customer data; PUBLIC demo profiles
only in GitHub workflows. No invented impact/version/remediation statements.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
from eclipse_reports.daily import _txt, _rect, _strip, C


def _page(prs: Presentation, heading: str, subtitle: str, page_no: int):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=C['bg']
    _rect(s,0,0,13.333,.13,'cyan')
    _txt(s,'ECLIPSE   /   EVIDENCE-BASED INTELLIGENCE',.65,.3,11,.32,size=10,color='muted',bold=True)
    _txt(s,heading,.65,.82,12.1,.68,size=27,bold=True)
    _txt(s,subtitle,.66,1.52,12,.31,size=11,color='muted')
    _rect(s,.68,7.1,12,.02,'line')
    _txt(s,'CVE  •  CISA KEV  •  FIRST EPSS  •  official advisory pointers',.72,7.16,11.6,.24,size=8,color='muted')
    _txt(s,str(page_no).zfill(2),12.27,7.14,.42,.26,size=9,color='muted')
    return s


def _label(slide,x,y,label,value,accent='blue'):
    _rect(slide,x,y,3.76,1.52,'white','line',radius=True)
    _rect(slide,x,y,.07,1.52,accent)
    _txt(slide,value,x+.29,y+.18,3.19,.57,size=31,color=accent,bold=True)
    _txt(slide,label,x+.3,y+.9,3.1,.37,size=13,bold=True)


def render_intelligence(briefing_path: Path, output_path: Path, *, max_details: int = 6, demo: bool = False) -> dict:
    raw=briefing_path.read_bytes()
    brief=json.loads(raw)
    if brief.get('schema_version')!=1 or not isinstance(brief.get('topics'),list):
        raise ValueError('Invalid multifeed briefing input')
    topics=brief['topics']
    prs=Presentation()
    prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    prs.core_properties.title='ECLIPSE / Public-source cross-feed threat intelligence'
    prs.core_properties.subject='Source-grounded automatic commentary, NOT verified customer exposure'
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=C['navy']
    _rect(s,.72,.9,.7,.11,'cyan')
    _txt(s,'PROJECT ECLIPSE',.78,1.23,11.7,.53,size=23,color='cyan',bold=True)
    _txt(s,'THREAT INTELLIGENCE\nDAILY BRIEFING',.78,2.0,11.67,1.5,size=35,color='white',bold=True)
    _txt(s,_strip(brief['customer_name'],70),.83,4.4,11.4,.56,size=25,color='white',bold=True)
    _txt(s,f"Source observation window: {brief['window']['start_utc']}  –  {brief['window']['end_utc']}",
         .83,5.16,11.7,.6,size=14,color='white')
    _txt(s,('SYNTHETIC DEMO — ALL OBSERVATIONS ARE FICTIONAL' if demo else 'Public-source analysis only • no customer environment assessed'),.83,6.4,11.5,.37,size=12,color='cyan')

    s=_page(prs,'エグゼクティブサマリー','指定監視条件に合う公開情報の観測結果',2)
    _label(s,.67,2.25,'抽出された事象数',str(len(topics)),'blue')
    _label(s,4.79,2.25,'KEVに新規登録を観測',str(sum(1 for x in topics if x['kev_added_in_window'])),'red')
    _label(s,8.91,2.25,'CVE原文取得済み',str(sum(1 for x in topics if x['technical_detail_status']=='official_CVE_record_fetched')),'green')
    _rect(s,.69,4.39,11.94,1.72,'white','line',radius=True)
    _txt(s,'読み方と制約',.94,4.63,11.2,.4,size=17,bold=True)
    desc=('この資料は指定期間にECLIPSEが観測した公式ソースの変化をまとめたもの。'
          '脆弱性の公開日・KEV登録日・攻撃開始日・初回検知時刻は異なる。'
          '対象製品の導入有無は照合していない。')
    _txt(s,desc,.94,5.12,11.15,.8,size=13)
    warnings=brief.get('source_warnings',[])
    _txt(s,('未収集ソース: '+_strip('; '.join(warnings),160)) if warnings else '監視ソースの状態: 取得済みデータのみを使用',
         .73,6.39,11.7,.4,size=10,color='red' if warnings else 'muted')

    # Render a bounded selection in PPTX; full matching topic set is retained
    # in the report JSON rather than building hundreds of unreadable slides.
    displayed=min(len(topics),40)
    for page in range(max(1,(displayed+7)//8)):
        subset=topics[page*8:min((page+1)*8,displayed)]
        s=_page(prs,'注目事象の一覧',f'観測単位で統合  /  スライド掲載 {displayed}/{len(topics)} 件  /  {page+1} ページ',len(prs.slides)+1)
        _rect(s,.73,2.0,11.9,.48,'navy2')
        _txt(s,'CVE / メーカー',.91,2.1,3.08,.3,size=11,color='white',bold=True)
        _txt(s,'観測された変化',4.07,2.1,3.49,.3,size=11,color='white',bold=True)
        _txt(s,'注目理由 / 注意点',7.6,2.1,4.84,.3,size=11,color='white',bold=True)
        if not subset:
            _txt(s,'この期間・条件に一致する観測事象はありません。監視対象外に脅威がないという意味ではありません。',
                 .95,3.0,11.2,.84,size=18,color='muted')
        for i,topic in enumerate(subset):
            y=2.63+i*.49
            _txt(s,_strip(topic['topic_id']+' / '+topic['vendor'],48),.91,y,2.98,.43,size=10,bold=True)
            _txt(s,_strip(', '.join(topic['event_types']),54),4.08,y,3.44,.43,size=9)
            _txt(s,_strip(topic['priority_label']+' : '+topic['why_now_ja'],74),7.6,y,4.82,.44,size=10)

    for ix, topic in enumerate(topics[:max_details],1):
        s=_page(prs,f"事象解説 {ix:02d}  /  {topic['topic_id']}",
                _strip(f"{topic['vendor']} / {topic['product']}  •  最終観測 {topic['last_observed_utc']}",150),len(prs.slides)+1)
        _rect(s,.71,2.02,7.5,4.78,'white','line',radius=True)
        _rect(s,8.38,2.02,4.24,4.78,'navy2',radius=True)
        _txt(s,'何が報告されているか',.97,2.25,6.86,.4,size=16,bold=True)
        _txt(s,_strip(topic['description_source_text'],380),.97,2.78,6.84,1.27,size=12)
        _txt(s,'なぜ今注目するか',.97,4.15,6.87,.38,size=16,bold=True)
        _txt(s,_strip(topic['why_now_ja'],320),.97,4.62,6.8,1.12,size=12)
        _txt(s,'監視・確認事項',.97,5.82,6.82,.34,size=15,bold=True)
        _txt(s,_strip(topic['next_checks_ja'],190),.97,6.17,6.83,.54,size=10)
        _txt(s,'出典付き評価',8.68,2.32,3.55,.4,size=17,color='cyan',bold=True)
        _txt(s,_strip(topic['priority_label'],45),8.69,2.95,3.55,.7,size=15,color='white',bold=True)
        _txt(s,'CVE原本の状態',8.68,3.85,3.62,.34,size=13,color='cyan',bold=True)
        _txt(s,'取得済み' if topic['technical_detail_status']=='official_CVE_record_fetched' else '未取得（詳細は未検証）',
             8.68,4.28,3.5,.6,size=12,color='white')
        _txt(s,'EPSS',8.68,5.07,3.5,.34,size=14,color='cyan',bold=True)
        x=topic.get('epss')
        _txt(s,(f"{x['score']:.4f} / {x['score_date']}" if x else '対象スコア未取得'),
             8.68,5.53,3.5,.4,size=14,color='white')
        _txt(s,'顧客資産への影響：未評価',8.69,6.28,3.48,.4,size=12,color='white')
        # The slide itself carries a clickable source, not just a JSON-only manifest.
        if topic.get('evidence'):
            first=topic['evidence'][0]['url']
            source_shape=_txt(s,'出典: '+_strip(first,125),.77,6.81,11.9,.25,size=9,color='blue')
            for paragraph in source_shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.hyperlink.address=first
    s=_page(prs,'出典・検証範囲','一次情報、変更時刻、未確定事項',len(prs.slides)+1)
    for i,(src,data) in enumerate(brief.get('source_coverage',{}).items()):
        _txt(s,f'{src}: '+('観測状態あり' if data.get('observed_state') else '未取得／未接続'),
             .94,2.03+i*.48,11.2,.36,size=14)
    _txt(s,'出典URLは個別事象のJSONマニフェストに全件記録。',.94,4.62,11.4,.4,size=13,bold=True)
    _txt(s,'CISA KEV: https://www.cisa.gov/known-exploited-vulnerabilities-catalog\n'
            'CVE: https://www.cve.org/  |  EPSS: https://www.first.org/epss/',
         .94,5.15,11.6,.81,size=12)
    _txt(s,'CVEの内容変更 ≠ 悪用確認。メーカーRSSの見出し ≠ 修正済みバージョンの証拠。\n'
            'この資料は自動生成の確認候補であり、顧客環境の影響判定ではない。',
         .94,6.1,11.5,.72,size=12,color='red')
    if demo:
        for slide in prs.slides:
            _txt(slide,'SYNTHETIC TEST DATA · NOT A LIVE ALERT',9.03,.27,3.64,.35,size=10,color='red',bold=True)
    output_path.parent.mkdir(parents=True,exist_ok=True)
    prs.save(str(output_path))
    manifest={'report_file':output_path.name,'briefing_sha256':hashlib.sha256(raw).hexdigest(),
              'selected_topics':[x['topic_id'] for x in topics],'topic_count':len(topics),
              'pptx_slide_count':len(prs.slides),'listed_in_slides_count':displayed,
              'omitted_from_list_slides_count':max(0,len(topics)-displayed),
              'synthetic_demo':demo,'sources':{
                  t['topic_id']:[e['url'] for e in t['evidence']] for t in topics},
              'caveats':brief.get('caveats',[])}
    output_path.with_suffix('.manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    return manifest


def main():
    p=argparse.ArgumentParser()
    p.add_argument('--briefing-json',type=Path,required=True)
    p.add_argument('--out',type=Path,required=True)
    p.add_argument('--max-details',type=int,default=6)
    p.add_argument('--demo',action='store_true')
    args=p.parse_args()
    if not 0<=args.max_details<=40: p.error('max-details must be 0..40')
    res=render_intelligence(args.briefing_json,args.out,max_details=args.max_details,demo=args.demo)
    print(json.dumps({k:v for k,v in res.items() if k not in ('sources','caveats')},ensure_ascii=False))

if __name__=='__main__':main()
