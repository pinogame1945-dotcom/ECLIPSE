"""Fully offline synthetic multisource tests: never labelled as real security alerts."""
import json
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from pptx import Presentation

from eclipse_l0.epss import SOURCE as EPSS_SOURCE
from eclipse_l1.timeline import compile_timeline, official_record_url
from eclipse_l1.briefing import make_briefing
from eclipse_l1.fetch_records import default_url, candidates, fetch_records
from eclipse_reports.intelligence import render_intelligence

NOW=datetime(2026,10,10,9,0,tzinfo=timezone.utc)
CVE='CVE-2026-12345'
OTHER='CVE-2026-12346'


def write_json(path,obj):
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(obj,ensure_ascii=False),encoding='utf-8')


class IntelligenceTests(unittest.TestCase):
    def setUp(self):
        tmp=tempfile.TemporaryDirectory();self.addCleanup(tmp.cleanup)
        self.root=Path(tmp.name)
        self.state=self.root/'data';self.state.mkdir()
        self.kev={
            'schema_version':1,'observed_utc':'2026-10-10T08:00:00+00:00','records':{
                CVE:{'cveID':CVE,'vendorProject':'Cisco','product':'TestOS',
                     'dateAdded':'2026-10-09','shortDescription':'Synthetic test-only auth bypass.',
                     'requiredAction':'Apply vendor supported updates.'}
            }
        }
        write_json(self.state/'kev-current.json',self.kev)
        write_json(self.state/'cve-delta-checkpoint.json',{'schema_version':1})
        self.kev_event={
            'event_id':'k1','kind':'added','cve_id':CVE,'source':'https://github.com/cisagov/kev-data',
            'observed_utc':'2026-10-10T07:00:00+00:00','date_added':'2026-10-09',
            'record':self.kev['records'][CVE], 'material_change':True,'changed_fields':[],
            'feed_sha256':'abcdef',
        }
        event_dir=self.state/'events';event_dir.mkdir()
        (event_dir/'2026-10.jsonl').write_text(json.dumps(self.kev_event)+'\n')
        self.cve_event={'kind':'cve_record_updated','cve_id':CVE,
            'observed_utc':'2026-10-10T07:30:00+00:00',
            'source_fetch_utc':'2026-10-10T07:29:00+00:00',
            'record_url':default_url(CVE),'delta_log_sha256':'digest'}
        d=self.state/'cve-observations';d.mkdir()
        (d/'2026-10.jsonl').write_text(json.dumps(self.cve_event)+'\n')
        self.out=self.root/'timeline.json'
        self.profile=self.root/'profile.json'
        write_json(self.profile,{'customer_name':'SYNTHETIC network demo','vendors':['Cisco'],'products':[],'include_all':False})
        self.all=self.root/'all.json'
        write_json(self.all,{'customer_name':'SYNTHETIC all sources','vendors':[],'products':[],'include_all':True})
        self.dirs=self.root/'cves';self.dirs.mkdir()
        self.cve_body={'cveMetadata':{'cveId':CVE,'state':'PUBLISHED'},
             'containers':{'cna':{'descriptions':[{'lang':'en','value':'Synthetic vulnerability in a demo component.'}],
             'affected':[{'vendor':'Cisco','product':'TestOS','versions':[{'version':'1.0','status':'affected'}]}]}}}
        write_json(self.dirs/(CVE+'.json'),self.cve_body)
        self.epss=self.root/'epss.json'
        write_json(self.epss,{'source':EPSS_SOURCE,'scores':{CVE:{'epss':0.52,'percentile':0.99}},
               'score_date':'2026-10-10'})

    def compile(self,**kwargs):
        return compile_timeline(self.state,self.out,hours=24,as_of=NOW,**kwargs)

    def briefing(self,profile=None,**kwargs):
        self.compile()
        return make_briefing(self.out,self.state/'kev-current.json',profile or self.profile,
                            self.root/'briefing.json',epss_file=self.epss,cve_dir=self.dirs,**kwargs)

    def test_two_sources_one_cve_with_provenance_and_pptx(self):
        tl=self.compile();self.assertEqual(tl['event_count'],2)
        r=make_briefing(self.out,self.state/'kev-current.json',self.profile,
                        self.root/'briefing.json',epss_file=self.epss,cve_dir=self.dirs)
        self.assertEqual(r['topics_count'],1)
        topic=r['topics'][0]
        self.assertEqual(topic['event_count'],2)
        self.assertTrue(topic['kev_added_in_window'])
        self.assertEqual(topic['epss']['score'],0.52)
        self.assertIn('CVE公式の技術説明',topic['description_source_text'])
        self.assertIn('CISA KEV',topic['why_now_ja'])
        self.assertEqual(len(topic['timeline']),2)
        self.assertTrue(topic['claims'])
        self.assertTrue(all(all(0 <= ref < len(topic['evidence']) for ref in claim['evidence_indices'])
                            for claim in topic['claims']))
        ppt=self.root/'briefing.pptx'
        manifest=render_intelligence(self.root/'briefing.json',ppt)
        self.assertGreaterEqual(len(Presentation(ppt).slides),5)
        self.assertIn(CVE,manifest['sources'])
        self.assertTrue((ppt.with_suffix('.manifest.json')).exists())

    def test_scoped_filter_keeps_unknown_vendor_as_explicit_gap(self):
        self.kev_event['cve_id']=OTHER
        self.kev_event['record']=dict(self.kev['records'][CVE],cveID=OTHER,vendorProject='')
        (self.state/'events/2026-10.jsonl').write_text(json.dumps(self.kev_event)+'\n')
        (self.state/'cve-observations/2026-10.jsonl').unlink()
        tl=self.compile()
        r=make_briefing(self.out,self.state/'kev-current.json',self.profile,self.root/'briefing.json')
        self.assertEqual(r['topics_count'],0)
        self.assertEqual(r['unclassified_excluded_count'],1)
        all_r=make_briefing(self.out,self.state/'kev-current.json',self.all,self.root/'all-brief.json')
        self.assertEqual(all_r['topics_count'],1)

    def test_no_baseline_fails_closed(self):
        (self.state/'kev-current.json').unlink()
        with self.assertRaisesRegex(ValueError,'baseline unavailable'):
            self.compile()

    def test_old_data_not_in_new_window(self):
        self.kev_event['observed_utc']='2026-10-05T01:00:00+00:00'
        self.cve_event['observed_utc']='2026-10-05T01:00:00+00:00'
        (self.state/'events/2026-10.jsonl').write_text(json.dumps(self.kev_event)+'\n')
        (self.state/'cve-observations/2026-10.jsonl').write_text(json.dumps(self.cve_event)+'\n')
        tl=self.compile();self.assertEqual(tl['event_count'],0)
        r=make_briefing(self.out,self.state/'kev-current.json',self.profile,self.root/'brief.json')
        self.assertEqual(r['topics_count'],0)
        self.assertGreaterEqual(len(Presentation(self._render(self.root/'brief.json')).slides),3)

    def _render(self,path):
        out=self.root/'daily.pptx';render_intelligence(path,out);return out

    def test_rejected_official_record_url(self):
        with self.assertRaisesRegex(ValueError,'Non-official'):
            official_record_url(CVE,'https://bad.example/CVE-2026-12345.json')
        with self.assertRaisesRegex(ValueError,'Non-official'):
            official_record_url(CVE,'https://raw.githubusercontent.com/CVEProject/cvelistV5/main/cves/2026/12xxx/'+CVE+'.json?secret=1')

    def test_advisory_without_cve_is_own_not_false_cve(self):
        advisory=self.root/'advisory.json'
        write_json(advisory,{
            'source_id':'testpub','baseline':False,'source_sha256':'abc','events':[
                {'vendor':'Cisco','title':'SYNTHETIC device update',
                 'description_excerpt':'No CVE number in this publication',
                 'link':'https://security.example.org/article','observed_utc':'2026-10-10T08:00:00Z',
                 'source_date':'Sat, 10 Oct 2026 08:00:00 GMT','cve_ids':[],'kind':'added'}]})
        tl=self.compile(advisory_files=[advisory]);self.assertEqual(tl['event_count'],3)
        r=make_briefing(self.out,self.state/'kev-current.json',self.profile,self.root/'brief.json',cve_dir=self.dirs)
        self.assertEqual(r['topics_count'],2)
        ad=[x for x in r['topics'] if x['cve_id'] is None][0]
        self.assertTrue(ad['topic_id'].startswith('ADVISORY-'))
        self.assertIn('公告',ad['why_now_ja'])

    def test_advisory_matching_cve_fuses_three_sources(self):
        advisory=self.root/'advisory.json'
        write_json(advisory,{'source_id':'testpub','baseline':False,'source_sha256':'abc','events':[
            {'vendor':'Cisco','title':'SYNTHETIC update '+CVE,'description_excerpt':'Summary',
             'link':'https://security.example.org/article','observed_utc':'2026-10-10T08:00:00Z',
             'source_date':'Sat, 10 Oct 2026 08:00:00 GMT','cve_ids':[CVE],'kind':'added'}]})
        self.compile(advisory_files=[advisory])
        out=make_briefing(self.out,self.state/'kev-current.json',self.profile,self.root/'brief.json')
        self.assertEqual(out['topics_count'],1)
        self.assertEqual(out['topics'][0]['event_count'],3)

    def test_cve_record_mismatch_rejected(self):
        self.compile()
        write_json(self.dirs/(CVE+'.json'),dict(self.cve_body,cveMetadata={'cveId':OTHER}))
        with self.assertRaisesRegex(ValueError,'ID mismatch'):
            make_briefing(self.out,self.state/'kev-current.json',self.profile,self.root/'brief.json',cve_dir=self.dirs)

    def test_bounded_cve_fetch_candidates_official(self):
        self.compile()
        doc=json.loads(self.out.read_text())
        selected=candidates(doc,2)
        self.assertEqual(selected[0][0],CVE)
        self.assertEqual(selected[0][1],default_url(CVE))
        with self.assertRaises(ValueError):
            candidates(doc,100)

    def test_provenance_preserves_unavailable_sources(self):
        (self.state/'cve-delta-checkpoint.json').unlink()
        tl=self.compile()
        self.assertFalse(tl['coverage']['cve_delta']['observed_state'])
        self.assertTrue(any('cve_delta' in x for x in tl['warnings']))

if __name__=='__main__':unittest.main()
