"""Synthetic-only, offline product portfolio to CVE candidate report regression."""
import csv
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

from pptx import Presentation
from eclipse_l0.core import connect, _upsert_cve
from eclipse_l0.warehouse import _meta, _schema
from eclipse_reports.customer_products import build_report, load_inventory
from eclipse_reports.customer_product_slides import render_report


class CustomerProductReportTests(unittest.TestCase):
    def setUp(self):
        t=tempfile.TemporaryDirectory()
        self.addCleanup(t.cleanup)
        self.root=Path(t.name)
        self.db=self.root/"warehouse.sqlite3"
        self.inventory=self.root/"synthetic-inventory.json"
        self.out=self.root/"output"
        self.con=connect(self.db)
        _schema(self.con)
        for k,v in (
            ("source_index_sha256","a"*64),("source_archive_sha256","b"*64),
            ("cve_full_ingest_verified",True),("kev_source_sha256","c"*64),
            ("epss_source_sha256","d"*64),("epss_score_date","2026-10-10T00:00:00Z")
        ):_meta(self.con,k,v)
        self.con.commit()
        self.addCleanup(self.con.close)
        self.write_inventory({
            "schema_version":1,"customer_name":"SYNTHETIC DEMO CUSTOMER",
            "products":[
                {"id":"fw-a","vendor":"Fortinet","product":"FortiOS",
                 "installed_version":"7.2.6"},
                {"id":"switch-b","vendor":"Cisco","product":"IOS XE",
                 "installed_version":"17.3.1"},
            ]
        })

    def write_inventory(self,d):
        self.inventory.write_text(json.dumps(d),encoding="utf-8")

    def add_cve(self,num,*,vendor="Fortinet",product="FortiOS",published="2026-10-09",
                updated="2026-10-09",cvss=8.7,epss=.2,kev=False,
                state="PUBLISHED",versions=None,status=None,container="cna"):
        cve=f"CVE-2026-{num:04d}"
        doc={"cveMetadata":{"cveId":cve,"state":state,
              "datePublished":published+"T00:00:00Z",
              "dateUpdated":updated+"T00:00:00Z"},"containers":{"cna":{}}}
        raw=json.dumps(doc).encode()
        _upsert_cve(self.con,doc,raw,"FICTIONAL_FIXTURE")
        self.con.execute("""INSERT INTO warehouse_summary
            (cve_id,description_en,cvss_score,cvss_kind)
            VALUES (?,?,?,?)""",(cve,"Fictional vulnerability test case",cvss,
                                 "cvssV3_1" if cvss is not None else None))
        if vendor is not None:
            self.con.execute("""INSERT INTO warehouse_affected
               (cve_id,container,ordinal,vendor,product,vendor_norm,product_norm,default_status,versions_json)
               VALUES (?,?,?,?,?,?,?,?,?)""",
               (cve,container,0,vendor,product,vendor.casefold(),product.casefold(),status,
                json.dumps(versions) if versions is not None else None))
        if epss is not None:self.con.execute("INSERT INTO warehouse_epss VALUES (?,?,?)",(cve,epss,.88))
        if kev:
            h=self.con.execute("SELECT hash FROM cve_current WHERE cve_id=?",(cve,)).fetchone()[0]
            self.con.execute("INSERT INTO kev_current VALUES (?,?,?,?,?)",(cve,h,"2026-10-08",1,"2026-10-10T00:00:00Z"))
        self.con.commit()
        return cve

    def build(self,*,days=90,mode="recent",preview=120):
        total=self.con.execute("SELECT COUNT(*) FROM cve_current").fetchone()[0]
        return build_report(self.db,self.inventory,self.out,as_of=date(2026,10,10),
                            days=days,mode=mode,expected_count=total,preview=preview,demo=True)

    def test_exact_vendor_product_no_substring_match(self):
        self.add_cve(1001,vendor="FORTINET",product="FORTIOS")
        self.add_cve(1002,vendor="notFortinet",product="FortiOS")
        self.add_cve(1003,vendor="Fortinet",product="FortiManager")
        report=self.build()
        self.assertEqual(report["summary"]["product_cve_rows_in_report"],1)
        self.assertEqual(report["preview"][0]["cve_id"],"CVE-2026-1001")
        self.assertEqual(report["preview"][0]["customer_version_impact"],"NOT_ASSESSED")

    def test_kevs_included_even_when_older_than_recent_window(self):
        self.add_cve(1001,published="2020-01-01",updated="2020-01-01",kev=True,epss=.001)
        self.add_cve(1002,published="2020-01-01",updated="2020-01-01",cvss=10)
        r=self.build()
        self.assertEqual(r["summary"]["product_cve_rows_in_report"],1)
        self.assertEqual(r["preview"][0]["review_queue"],"KEV_LISTED_REVIEW")
        self.assertEqual(r["summary"]["published_product_cve_pairs_all_history"],2)
        self.assertEqual(self.build(mode="all")["summary"]["product_cve_rows_in_report"],2)

    def test_publisher_unaffected_version_never_certifies_customer_impact(self):
        self.add_cve(1001,versions=[{"version":"7.2.6","status":"unaffected"}],status="unaffected")
        r=self.build()
        entry=r["preview"][0]
        self.assertEqual(entry["publisher_version_statements"][0]["default_status"],"unaffected")
        self.assertEqual(entry["customer_version_impact"],"NOT_ASSESSED")
        self.assertEqual(entry["installed_version"],"7.2.6")
        self.assertEqual(entry["source_match_kind"],"PUBLISHER_AFFECTED_ARRAY_PRODUCT_LABEL")

    def test_aliases_explicit_only_and_duplicate_alias_ignored(self):
        self.write_inventory({"schema_version":1,"customer_name":"SYNTHETIC DEMO",
            "products":[{"id":"fw","vendor":"Fortinet","product":"FortiOS",
                         "aliases":[{"vendor":"Fortinet","product":"FortiOS"},
                                    {"vendor":"Fortinet","product":"FortiOS Operating System"}]}]})
        self.add_cve(1001,product="FortiOS Operating System")
        self.add_cve(1002)
        r=self.build()
        self.assertEqual(r["summary"]["product_cve_rows_in_report"],2)
        self.assertEqual(r["summary"]["monitored_inventory_entries"],1)

    def test_incomplete_product_names_do_not_generate_unfiltered_report(self):
        for d in [
            {"schema_version":1,"customer_name":"x","products":[]},
            {"schema_version":1,"customer_name":"x","products":[{"vendor":"Cisco"}]},
            {"schema_version":1,"customer_name":"x","products":[{"vendor":"Cisco","product":"IOS XE","id":"a"},
               {"vendor":"Cisco","product":"IOS XE","id":"a"}]},
            {"schema_version":1,"customer_name":"x","products":[{"vendor":"Cisco","product":"IOS XE",
                      "aliases":[{"vendor":"Cisco","product":""}]}]},
        ]:
            self.write_inventory(d)
            with self.assertRaises(ValueError):
                load_inventory(self.inventory)

    def test_empty_match_is_unknown_not_safe(self):
        self.add_cve(1001,vendor=None)
        r=self.build()
        self.assertEqual(r["summary"]["product_cve_rows_in_report"],0)
        self.assertIn("do NOT establish safety"," ".join(r["limitations"]))
        self.assertEqual(r["preview"],[])

    def test_preview_limit_does_not_truncate_full_csv(self):
        for n in range(1001,1016):
            self.add_cve(n,epss=n/10000)
        r=self.build(preview=3)
        self.assertEqual(r["summary"]["product_cve_rows_in_report"],15)
        self.assertEqual(len(r["preview"]),3)
        self.assertEqual(r["preview_omitted_rows"],12)
        with (self.out/"candidate-vulnerabilities.csv").open(encoding="utf-8-sig",newline="") as f:
            rows=list(csv.DictReader(f))
        self.assertEqual(len(rows),15)
        self.assertTrue(all(x["customer_version_impact"]=="NOT_ASSESSED" for x in rows))

    def test_csv_formula_injection_is_escaped(self):
        self.write_inventory({"schema_version":1,"customer_name":"SYNTHETIC",
            "products":[{"vendor":"Fortinet","product":"FortiOS","installed_version":"=HYPERLINK(\"x\")"}]})
        self.add_cve(1001)
        self.build()
        with (self.out/"candidate-vulnerabilities.csv").open(encoding="utf-8-sig") as f:
            cells=list(csv.DictReader(f))
        self.assertTrue(cells[0]["installed_version"].startswith("'="))

    def test_pptx_renders_all_products_with_factual_caveats(self):
        self.add_cve(1001,kev=True)
        report=self.build()
        pptx=self.out/"synthetic-report.pptx"
        m=render_report(self.out/"customer-product-report.json",pptx,max_details=1)
        prs=Presentation(pptx)
        text=" ".join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
        self.assertIn("製品別",text)
        self.assertIn("FortiOS",text)
        self.assertIn("CVE-2026-1001",text)
        self.assertIn("未判定",text)
        self.assertGreaterEqual(m["slides"],5)
        self.assertEqual(m["reported_product_cve_rows"],1)

    def test_source_pin_and_cve_count_are_checked(self):
        self.add_cve(1001)
        from eclipse_l0.warehouse import SOURCE_INDEX_SHA256,SOURCE_SHA256
        with self.assertRaisesRegex(ValueError,"pinned full snapshot"):
            build_report(self.db,self.inventory,self.out,as_of=date(2026,10,10),expected_count=403696,demo=True)
        _meta(self.con,"source_index_sha256",SOURCE_INDEX_SHA256)
        _meta(self.con,"source_archive_sha256",SOURCE_SHA256)
        self.con.commit()
        with self.assertRaisesRegex(ValueError,"coverage"):
            build_report(self.db,self.inventory,self.out,as_of=date(2026,10,10),expected_count=403696,demo=True)

    def test_unpublished_and_adp_duplicates(self):
        self.add_cve(1001,state="REJECTED")
        self.add_cve(1002)
        self.con.execute("""INSERT INTO warehouse_affected
           (cve_id,container,ordinal,vendor,product,vendor_norm,product_norm)
           VALUES (?,?,?,?,?,?,?)""",("CVE-2026-1002","adp:0",0,"Fortinet","FortiOS","fortinet","fortios"))
        self.con.commit()
        r=self.build()
        self.assertEqual(r["summary"]["product_cve_rows_in_report"],1)
        self.assertEqual(len(r["preview"][0]["publisher_version_statements"]),2)


if __name__=="__main__":unittest.main()
