import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from pptx import Presentation
from eclipse_reports.daily import _select, render

ROOT = Path(__file__).resolve().parents[1]
FEED = json.loads((ROOT / 'examples/SYNTHETIC-kev-sample.json').read_text(encoding='utf8'))

class DailyReportTests(unittest.TestCase):
    def test_vendor_selection(self):
        cfg = {'customer_name': 'x', 'vendors': ['Cisco', 'Fortinet'], 'products': []}
        rows = _select(FEED, cfg, date(2026,10,9), date(2026,10,9))
        self.assertEqual([r['cveID'] for r in rows], ['CVE-2099-00002', 'CVE-2099-00001'])
    def test_empty_matches_are_valid(self):
        cfg = {'customer_name': 'x', 'vendors': ['NoSuchVendor'], 'products': []}
        self.assertEqual(_select(FEED, cfg, date(2026,10,9), date(2026,10,9)), [])
    def test_no_filter_without_opt_in_is_rejected(self):
        with self.assertRaises(ValueError):
            _select(FEED, {'customer_name':'x'}, date(2026,10,9), date(2026,10,9))
        with self.assertRaises(ValueError):
            _select(FEED, {'customer_name':'x'}, date(1999,1,1), date(1999,1,1))
    def test_product_intersection(self):
        cfg = {'customer_name': 'x', 'vendors': ['Cisco'], 'products':['router']}
        rows = _select(FEED, cfg, date(2026,10,1), date(2026,10,9))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['cveID'], 'CVE-2099-00004')
    def test_powerpoint_and_manifest(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / 'report.pptx'
            manifest = render(ROOT / 'examples/SYNTHETIC-kev-sample.json', ROOT / 'config/customer-network-example.json', output,
                             date(2026,10,9), date(2026,10,9), demo=True)
            self.assertEqual(manifest['selected_count'],2)
            ppt = Presentation(output)
            self.assertEqual(len(ppt.slides), 6)
            self.assertIn('2099-00002', ' '.join(sh.text for s in ppt.slides for sh in s.shapes if sh.has_text_frame))
            self.assertTrue(output.with_suffix('.manifest.json').exists())

if __name__ == '__main__': unittest.main()

class ProductionFeedGuardTests(unittest.TestCase):
    def test_synthetic_feed_rejected_when_not_demo(self):
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaises(ValueError):
                render(ROOT / 'examples/SYNTHETIC-kev-sample.json',
                       ROOT / 'config/customer-network-example.json',
                       Path(folder) / 'no.pptx', date(2026,10,9), date(2026,10,9))
