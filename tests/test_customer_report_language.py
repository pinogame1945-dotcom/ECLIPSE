"""Client slide QA: Japanese summaries and readable version ranges only."""
import json
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from eclipse_reports.customer_product_slides import (clean_version, jpversion,
   qa_ja_summary, reason, render_report)


class CustomerSlideLanguageTests(unittest.TestCase):
    def test_no_internal_machine_labels_or_json_verbatim(self):
        report={"schema_version":1,
            "classification":"CUSTOMER_PRODUCT_REFERENCE_REPORT_NOT_VERIFIED_ASSET_EXPOSURE",
            "customer_name":"SYNTHETIC CLIENT",
            "synthetic_demo":True,"as_of":"2026-10-10",
            "mode":"recent","lookback_start_inclusive":"2026-07-12",
            "summary":{"monitored_inventory_entries":1,
                "product_cve_rows_in_report":1,"kev_listed_product_cve_rows":1},
            "products":[{"vendor":"Fortinet","product":"FortiOS","included_product_cve_rows":1,
                "kev_listed_included":1}],
            "preview":[{"cve_id":"CVE-2099-0001","vendor":"Fortinet","product":"FortiOS",
                "kev_in_snapshot":True,"review_queue":"KEV_LISTED_REVIEW","cvss_score":9.7,
                "publisher_description":"A hypothetical vulnerability in FooOS that shouldn't go to the client",
                "publisher_version_statements":[{"container":"cna","versions_raw":[
                    {"version":"1.0.0","lessThanOrEqual":"1.0.4","status":"affected","versionType":"semver"}]}]}]
        }
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            rp=root/"report.json";rp.write_text(json.dumps(report),encoding="utf8")
            summaries=root/"approved.json"
            summaries.write_text(json.dumps({"CVE-2099-0001":"テスト用製品の認証処理に問題があります。"}),encoding="utf8")
            m=render_report(rp,root/"out.pptx",briefs_ja=summaries,max_details=1)
            prs=Presentation(root/"out.pptx")
            text=" ".join(sh.text for slide in prs.slides for sh in slide.shapes if sh.has_text_frame)
            for token in ("KEV_LISTED_REVIEW","NOT_ASSESSED","EXAMPLE-ONLY","versions_raw",
                          "affected[]","A hypothetical vulnerability","SYNTHETIC CLIENT"):
                self.assertNotIn(token,text)
            self.assertIn("テスト用製品の認証処理に問題",text)
            self.assertIn("1.0.0 ～ 1.0.4",text)
            self.assertIn("未判定",text)
            self.assertGreaterEqual(m["slides"],5)

    def test_versions_prefer_cna_and_do_not_emit_json(self):
        example=[{"container":"adp:0","versions_raw":[
            {"version":"7.2.0","lessThan":"7.2.7","status":"affected"}]},
            {"container":"cna","versions_raw":[
             {"version":"7.2.0","lessThanOrEqual":"7.2.6","status":"affected"},
             {"version":"7.4.0","lessThanOrEqual":"7.4.2","status":"affected"},
             {"version":"8.0.1","status":"unaffected"}]}]
        self.assertEqual(jpversion(example),"7.2.0 ～ 7.2.6、7.4.0 ～ 7.4.2")

    def test_summary_rejects_english_and_fake_missing_version(self):
        self.assertEqual(clean_version("EXAMPLE-ONLY-UNKNOWN"),"未確認")
        self.assertEqual(clean_version(None),"未確認")
        with self.assertRaisesRegex(ValueError,"Japanese summary"):
            qa_ja_summary({"CVE-2099-0001":"This is an English summary"})
        self.assertEqual(reason({"kev_in_snapshot":False,"review_queue":"RECENT_CVE_PUBLICATION_OR_UPDATE"}),
                         "最近公開・更新")


if __name__=="__main__":unittest.main()
